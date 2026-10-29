from pathlib import Path

def _rows(tasks):
    return [[t['title'],t['person_name'],t['task_date'],t['start_time']+' - '+t['end_time'],str(t['priority'])+'/10',t['status'],t['details'] or ''] for t in tasks]

def export_xlsx(tasks,path,start,end):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    wb=Workbook(); ws=wb.active; ws.title='المهام'; ws.sheet_view.rightToLeft=True
    ws.append(['تقرير المهام','الفترة: '+start+' إلى '+end]); ws.merge_cells('A1:G1'); ws['A1'].font=Font(name='Arial',size=16,bold=True); ws['A1'].alignment=Alignment(horizontal='center')
    headers=['المهمة','المسؤول','التاريخ','الوقت','الأولوية','الحالة','التفاصيل']; ws.append(headers)
    for c in ws[2]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='24476B'); c.alignment=Alignment(horizontal='center',vertical='center')
    for row in _rows(tasks): ws.append(row)
    widths=[32,20,16,20,12,18,48]
    for i,w in enumerate(widths,1): ws.column_dimensions[chr(64+i)].width=w
    for row in ws.iter_rows(min_row=3):
        for c in row: c.alignment=Alignment(horizontal='right',vertical='top',wrap_text=True)
    ws.freeze_panes='A3'; ws.auto_filter.ref=f'A2:G{ws.max_row}'; wb.save(path)

def export_docx(tasks,path,start,end):
    from docx import Document
    from docx.shared import Inches,Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    doc=Document(); sec=doc.sections[0]; sec.right_to_left=True
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('تقرير مهام الفريق'); r.bold=True; r.font.size=Pt(18)
    p=doc.add_paragraph(f'الفترة: {start} إلى {end}'); p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    headers=['المهمة','المسؤول','التاريخ','الوقت','الأولوية','الحالة','التفاصيل']; table=doc.add_table(rows=1,cols=len(headers)); table.style='Table Grid'; table.alignment=WD_TABLE_ALIGNMENT.CENTER
    for c,h in zip(table.rows[0].cells,headers): c.text=h; c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row in _rows(tasks):
        cells=table.add_row().cells
        for c,v in zip(cells,row): c.text=str(v); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    doc.add_paragraph(f'إجمالي المهام: {len(tasks)}'); doc.save(path)

def export_pdf(tasks,path,start,end):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_RIGHT, TA_CENTER
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    # Register a Windows Arabic font if available. ReportLab's basic fonts do not shape Arabic.
    candidates=[Path('C:/Windows/Fonts/arial.ttf'),Path('C:/Windows/Fonts/tahoma.ttf'),Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')]
    font='Helvetica'
    for f in candidates:
        if f.exists():
            try: pdfmetrics.registerFont(TTFont('ArabicFont',str(f))); font='ArabicFont'; break
            except Exception: pass
    doc=SimpleDocTemplate(path,pagesize=landscape(A4),rightMargin=24,leftMargin=24,topMargin=24,bottomMargin=24)
    title=ParagraphStyle('ArabicTitle',fontName=font,fontSize=16,leading=22,alignment=TA_CENTER,spaceAfter=8)
    normal=ParagraphStyle('ArabicBody',fontName=font,fontSize=8,leading=11,alignment=TA_RIGHT)
    story=[Paragraph('تقرير مهام الفريق',title),Paragraph(f'الفترة: {start} إلى {end} | إجمالي المهام: {len(tasks)}',normal),Spacer(1,10)]
    headers=['المهمة','المسؤول','التاريخ','الوقت','الأولوية','الحالة','التفاصيل']; data=[[Paragraph(x,normal) for x in headers]]
    for row in _rows(tasks): data.append([Paragraph(str(x).replace('&','&amp;').replace('<','&lt;').replace('>','&gt;').replace('\n','<br/>'),normal) for x in row])
    table=Table(data,repeatRows=1,colWidths=[155,90,70,75,55,75,220]); table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#24476B')),('TEXTCOLOR',(0,0),(-1,0),colors.white),('GRID',(0,0),(-1,-1),0.4,colors.grey),('VALIGN',(0,0),(-1,-1),'TOP'),('ALIGN',(0,0),(-1,-1),'RIGHT'),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#F2F5F8')]),('LEFTPADDING',(0,0),(-1,-1),5),('RIGHTPADDING',(0,0),(-1,-1),5),('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5)])); story.append(table); doc.build(story)


def export_csv(tasks, path, start='', end=''):
    """UTF-8 (with BOM so Excel shows Arabic correctly) CSV of the given tasks."""
    import csv
    with open(path, 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.writer(f)
        w.writerow(['المهمة', 'المسؤول', 'الفئة', 'التاريخ', 'البداية', 'النهاية', 'الأولوية', 'الحالة', 'التذكير (دقيقة)', 'التفاصيل'])
        for t in tasks:
            w.writerow([t['title'], t['person_name'], (t['category'] if 'category' in t.keys() else '') or '', t['task_date'], t['start_time'],
                        t['end_time'], t['priority'], t['status'], t['reminder_minutes'], t['details'] or ''])
